#!/usr/bin/env python3
"""Convert DOCX to Markdown with extracted images at original positions.

Text, headings, and tables are extracted directly via python-docx (lossless).
Embedded images larger than the threshold are saved to disk and referenced
via standard Markdown ![](path) at their original position. The calling agent
(Claude Code) describes the images by reading them with the Read tool — no
API key required.

For backend / non-interactive use, pass --api-key to have the script call
Claude Vision itself and inline the descriptions.

Usage:
  # Agent mode (default, zero config):
  python docx_to_md.py --input file.docx --output ./out/

  # Standalone mode (script calls Vision directly):
  python docx_to_md.py --input file.docx --output ./out/ --api-key sk-ant-...

  # Pure text, skip image extraction entirely:
  python docx_to_md.py --input file.docx --output ./out/ --no-images
"""
import argparse
import base64
import os
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("ERROR: pip install python-docx")

VISION_PROMPT = (
    "请描述这张图片的完整内容。"
    "如果是流程图或架构图，描述各节点名称和它们之间的连接关系（箭头方向、流程顺序）；"
    "如果是图表或数据表，提取关键数字、坐标轴标签和趋势；"
    "如果是界面截图，描述关键功能区域和显示的数据；"
    "如果是对比图，说明左右/上下各自代表什么及其差异。"
    "直接输出 Markdown 格式内容，不要添加前言或后记。"
)

# Bytes that identify raster image formats
RASTER_SIGNATURES = {
    b"\xff\xd8\xff": ("image/jpeg", "jpg"),
    b"\x89PNG": ("image/png", "png"),
    b"GIF8": ("image/gif", "gif"),
    b"BM": ("image/bmp", "bmp"),
    b"RIFF": ("image/webp", "webp"),  # RIFF....WEBP
}


def detect_format(blob: bytes) -> tuple[str, str] | None:
    """Return (media_type, file_extension) or None for unsupported formats."""
    for sig, info in RASTER_SIGNATURES.items():
        if blob[: len(sig)] == sig:
            return info
    return None


def iter_block_items(doc: Document):
    """Yield Paragraph and Table objects in document body order."""
    body = doc.element.body
    for child in body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def heading_prefix(para: Paragraph) -> str:
    style = para.style.name if para.style else ""
    if style.startswith("Heading"):
        try:
            level = int(style.split()[-1])
            return "#" * min(level, 6) + " "
        except ValueError:
            pass
    return ""


def para_to_md(para: Paragraph) -> str:
    prefix = heading_prefix(para)
    text = para.text.strip()
    if not text:
        return ""
    return prefix + text


def table_to_md(table: Table) -> str:
    rows = []
    for i, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
        if i == 0:
            rows.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return "\n".join(rows)


def extract_images_from_para(para: Paragraph, doc: Document) -> list[tuple[str, bytes]]:
    """Return list of (relationship_id, blob) for all inline images in the paragraph."""
    images = []
    for blip in para._element.findall(".//" + qn("a:blip")):
        r_id = blip.get(qn("r:embed"))
        if not r_id:
            continue
        try:
            part = doc.part.related_parts[r_id]
            images.append((r_id, part.blob))
        except (KeyError, AttributeError):
            pass
    return images


def call_vision(blob: bytes, media_type: str, model: str, client) -> str:
    b64 = base64.standard_b64encode(blob).decode()
    resp = client.messages.create(
        model=model,
        max_tokens=1024,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": media_type,
                            "data": b64,
                        },
                    },
                    {"type": "text", "text": VISION_PROMPT},
                ],
            }
        ],
    )
    return resp.content[0].text.strip()


def convert(
    docx_path: Path,
    output_dir: Path,
    extract_images: bool,
    large_image_kb: int,
    max_images: int,
    standalone_client=None,
    model: str = "claude-haiku-4-5-20251001",
) -> tuple[Path, int]:
    """
    Returns (output_path, image_count).
    If standalone_client is provided, descriptions are inlined.
    Otherwise images are saved to disk and referenced via ![](path) placeholders.
    """
    doc = Document(str(docx_path))
    threshold = large_image_kb * 1024
    stem = docx_path.stem
    imgs_dir = output_dir / stem / "imgs"

    lines: list[str] = []
    images_processed = 0

    for block in iter_block_items(doc):
        if isinstance(block, Table):
            md = table_to_md(block)
            if md:
                lines.append(md)
                lines.append("")
        elif isinstance(block, Paragraph):
            if extract_images and images_processed < max_images:
                for _r_id, blob in extract_images_from_para(block, doc):
                    if len(blob) < threshold:
                        continue
                    fmt = detect_format(blob)
                    if not fmt:
                        continue  # skip EMF/WMF/vector
                    media_type, ext = fmt
                    images_processed += 1
                    if standalone_client is not None:
                        try:
                            description = call_vision(blob, media_type, model, standalone_client)
                            lines.append(f"\n> **[图片]**\n>\n> {description}\n")
                        except Exception as e:
                            lines.append(f"\n> *[图片处理失败: {e}]*\n")
                    else:
                        imgs_dir.mkdir(parents=True, exist_ok=True)
                        img_name = f"img_{images_processed:03d}.{ext}"
                        (imgs_dir / img_name).write_bytes(blob)
                        lines.append(f"\n![]({stem}/imgs/{img_name})\n")
                    if images_processed >= max_images:
                        lines.append(
                            f"\n> *（达到单文档最大图片处理数 {max_images}，"
                            "后续图片已跳过）*\n"
                        )

            text = para_to_md(block)
            if text:
                lines.append(text)

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / (stem + ".md")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path, images_processed


def main():
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("--input", required=True, help="DOCX file or directory of DOCX files")
    ap.add_argument("--output", required=True, help="Output directory")
    ap.add_argument(
        "--large-image-kb",
        type=int,
        default=30,
        help="Images larger than this (KB) are extracted/described (default: 30)",
    )
    ap.add_argument(
        "--max-images",
        type=int,
        default=50,
        help="Max images to process per document (cost guard, default: 50)",
    )
    ap.add_argument(
        "--no-images",
        action="store_true",
        help="Skip image extraction entirely (text-only output)",
    )
    ap.add_argument(
        "--api-key",
        default=None,
        help="Anthropic API key. If provided, the script calls Vision itself "
             "and inlines descriptions (standalone mode for backend / cron use). "
             "Default: extract images to disk and emit ![](...) placeholders "
             "for the calling agent to fill in.",
    )
    ap.add_argument(
        "--model",
        default=os.environ.get("DOCS_TO_WIKI_MODEL", "claude-haiku-4-5-20251001"),
        help="Vision model when --api-key is set (env: DOCS_TO_WIKI_MODEL)",
    )
    args = ap.parse_args()

    extract_images = not args.no_images
    input_path = Path(args.input).resolve()
    output_dir = Path(args.output).resolve()

    standalone_client = None
    if args.api_key:
        try:
            import anthropic
        except ImportError:
            sys.exit("ERROR: --api-key requires `pip install anthropic`")
        standalone_client = anthropic.Anthropic(api_key=args.api_key)

    docx_files = (
        sorted(input_path.glob("*.docx"))
        if input_path.is_dir()
        else [input_path]
    )

    if not docx_files:
        print("No DOCX files found.")
        return

    total_images = 0
    mode = "standalone" if standalone_client else ("agent" if extract_images else "text-only")
    print(f"Mode: {mode}")
    for i, docx_path in enumerate(docx_files, 1):
        print(f"[{i}/{len(docx_files)}] {docx_path.name} ...", end=" ", flush=True)
        _out, n_imgs = convert(
            docx_path,
            output_dir,
            extract_images=extract_images,
            large_image_kb=args.large_image_kb,
            max_images=args.max_images,
            standalone_client=standalone_client,
            model=args.model,
        )
        total_images += n_imgs
        verb = "described" if standalone_client else "extracted"
        print(f"done ({n_imgs} images {verb})")

    print(f"\nTotal: {len(docx_files)} docs, {total_images} images → {output_dir}")
    if mode == "agent" and total_images > 0:
        print("\nNext: ask Claude Code to fill in the ![](...) image placeholders.")


if __name__ == "__main__":
    main()
